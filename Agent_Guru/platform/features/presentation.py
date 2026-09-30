"""File-to-presentation generation for Agent Guru.

Turns an arbitrary document (a project file, a set of notes) into a
real, valid .pptx a person can open in PowerPoint - offline, no cloud
services and no API key required for the default path.

Three cooperating pieces, from simplest to richest:

1. Source extraction (_read_source_text): .txt/.md are read natively.
   .docx (python-docx) and .pdf (pypdf) are OPTIONAL soft dependencies
   - matching the rest of this codebase's pattern (see
   platform/features/web_search.py) - a missing library raises a
   PresentationError with an actionable install hint rather than a
   bare ImportError or a crash.
2. Outline building: a heuristic, dependency-free splitter
   (_heuristic_outline) always works and is markdown-header-aware, so
   a well-structured .md file produces a well-structured deck with no
   LLM involved at all. When an async LLM completion callable is
   available (the same shape Agent Guru's orchestrator already uses -
   see integration/orchestrator.py::_llm_complete, which builds a
   `CompletionRequest` and awaits `complete_fn(request)`), an optional
   richer outline (_llm_outline) asks the model to group the material
   into slide-sized chunks. If the model is unavailable, times out, or
   returns something that cannot be parsed as JSON, this silently
   falls back to the heuristic splitter rather than failing the whole
   request - generating *a* deck always succeeds for any non-empty
   source, LLM or not.
3. Rendering (render_pptx): turns a list of SlidePlan objects into a
   real python-pptx Presentation and saves it. This is the only part
   that touches PowerPoint machinery.

generate_presentation() is the single public entry point wiring 1-3
together end to end.

Wiring into the live pipeline (Round 10): this module IS now connected
to integration/orchestrator.py, via _detect_universal_features() /
_handle_presentation_request() there - the same "runs for every intent
type, before system_control/file_ops dispatch" hook that briefing/
system-status/dictation/clipboard/screen-reading already use.
wants_presentation() detects the request and extract_source_path()
(below) does the slot-filling: a quoted path or a bare filename ending
in a recognized extension, not full NLU resolution of "my notes" to an
actual file (a real limitation - a request with no identifiable file
extension falls through to the conversational LLM instead of guessing).
The orchestrator resolves the extracted path through
security.path_utils.safe_path() (the same sandbox
reasoning/tools/system/file_operations.py uses) before calling
`generate_presentation(source_path, complete_fn=getattr(self.llm_runtime,
"complete", None))`, and reports the resulting Path back to the user
(e.g. "Your presentation is ready at ..."). This was done as a small,
additive hook rather than an edit to the rule-based intent recognizer's
regex tables (platform/intent_recognizer.py) - this module was built
and tested standalone first, and is wired in now with a real executed
test covering both the wiring and the security boundary.

An additional, fully optional cloud path (generate_presentation_gamma)
talks to Gamma's public API (https://developers.gamma.app) instead of
rendering locally. It is opt-in only: it never activates unless the
caller passes an explicit api_key or has set GAMMA_API_KEY themselves
- offline-first means the local python-pptx path above is always what
runs by default.

Install (optional, only what you need):
    pip install python-pptx          # required for the local render path
    pip install python-docx          # only if reading .docx sources
    pip install pypdf                # only if reading .pdf sources
"""

from __future__ import annotations

import asyncio
import json
import logging
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from collections.abc import Match

logger = logging.getLogger(__name__)

_MAX_SLIDES_DEFAULT = 12
_MAX_BULLETS_PER_SLIDE = 6
_MAX_BULLET_CHARS = 200
_LLM_OUTLINE_TIMEOUT_S = 30.0
_LLM_SOURCE_CHAR_LIMIT = 12_000

# Cheap regex pre-filter an intent-recognition stage could use before
# doing any real work - mirrors web_search.py's needs_web_search().
# WIRED (Round 10): called from
# integration/orchestrator.py::_detect_universal_features(), the same
# "universal feature detection" hook that is_briefing_command()/
# is_dictation_command()/is_screen_read_command() already use - see
# _handle_presentation_request() there for the full wiring. That method
# runs for every intent type, before the rule-based intent recognizer's
# system_control/file_ops dispatch, so this needed no change to the
# 45KB rule-engine tables in platform/intent_recognizer.py.
# BUG FIX (Round 10): the original pattern only recognized "turn this into
# slides" / "turn it into slides" (an anaphoric reference to something
# already in context) but NOT "turn notes.txt into a presentation" or
# "convert report.docx to slides" -- both completely natural phrasings once
# this feature is actually reachable by a real spoken/typed source-file
# reference (found while writing the Round 10 end-to-end test: a quoted
# filename with "turn ... into" failed to trigger at all). Restructured as
# two alternatives: the original verb+object form, and a "turn/convert ...
# into/to ... <object>" form that doesn't care what's between the verb and
# "into"/"to" -- so any file reference works, not just "this"/"it".
_PRESENTATION_KEYWORDS = re.compile(
    r"\b(?:"
    r"(?:make|create|build|generate|prepare)\b.{0,60}\b(?:presentation|slides?|slide ?deck|deck|ppt|powerpoint|pptx)\b"
    r"|"
    r"(?:turn|convert)\b.{0,60}\b(?:into|to)\b.{0,20}\b(?:presentation|slides?|slide ?deck|deck|ppt|powerpoint|pptx)\b"
    r")",
    re.IGNORECASE,
)

# Matches an explicit source-file reference in the transcript: a quoted
# path, or a bare filename/path ending in a recognized extension. This is
# intentionally simple slot-filling (not full NLU) - it resolves "make a
# presentation from notes.txt" or "turn ~/Documents/report.docx into
# slides" but NOT "make a presentation from my notes" (no file extension
# to anchor on). The latter is left to fall through to the conversational
# LLM, which can ask a clarifying question, rather than guessing a file.
_SOURCE_FILE_EXTENSIONS = r"txt|md|markdown|rst|log|docx|pdf"
_QUOTED_PATH_RE = re.compile(r"""["']([^"']+\.(?:""" + _SOURCE_FILE_EXTENSIONS + r"""))["']""", re.IGNORECASE)
_BARE_PATH_RE = re.compile(
    r"\b([\w~][\w./\\-]*\.(?:" + _SOURCE_FILE_EXTENSIONS + r"))\b",
    re.IGNORECASE,
)


def wants_presentation(query: str) -> bool:
    """Return True if the query is asking to turn something into slides."""
    return bool(_PRESENTATION_KEYWORDS.search(query))


def extract_source_path(query: str) -> str | None:
    """Best-effort extraction of a source file path from a presentation request.

    Checks a quoted path first (least ambiguous), then a bare filename/path
    token ending in a known extension. Returns None if neither is found -
    callers should treat that as "couldn't identify a file" rather than
    guessing, since generating a presentation from the wrong file silently
    would be worse than asking the user which file they meant.
    """
    match = _QUOTED_PATH_RE.search(query)
    if match:
        return match.group(1)
    match = _BARE_PATH_RE.search(query)
    if match:
        return match.group(1)
    return None


@dataclass
class SlidePlan:
    """One planned slide: a title and its bullet points."""

    title: str
    bullets: list[str] = field(default_factory=list)


class PresentationError(Exception):
    """Raised when a presentation cannot be generated (bad source, missing
    optional dependency, empty content, etc.) - always carries an
    actionable message rather than surfacing a raw library exception.
    """


# ---------------------------------------------------------------------------
# 1. Source extraction
# ---------------------------------------------------------------------------


async def _read_source_text(path: Path) -> str:
    """Extract plain text from a source file.

    .txt/.md/.markdown are read directly. .docx and .pdf are supported
    only if python-docx / pypdf are installed; any other extension is
    given a best-effort plain-text read (covers .rst, .log, source
    code, etc.) rather than being refused outright.
    """
    suffix = path.suffix.lower()
    loop = asyncio.get_running_loop()

    if suffix == ".docx":
        return await _read_docx(path, loop)
    if suffix == ".pdf":
        return await _read_pdf(path, loop)

    try:
        return await loop.run_in_executor(
            None,
            lambda: path.read_text(encoding="utf-8", errors="replace"),
        )
    except OSError as exc:
        msg = f"Could not read {path}: {exc}"
        raise PresentationError(msg) from exc


async def _read_docx(path: Path, loop: asyncio.AbstractEventLoop) -> str:
    try:
        import docx  # type: ignore[import]  # python-docx
    except ImportError as exc:
        msg = "Reading .docx files requires python-docx. Install with: pip install python-docx"
        raise PresentationError(msg) from exc

    def _extract() -> str:
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text.strip())
        return "\n\n".join(parts)

    try:
        return await loop.run_in_executor(None, _extract)
    except PresentationError:
        raise
    except Exception as exc:
        msg = f"Failed to parse {path.name} as a .docx file: {exc}"
        raise PresentationError(msg) from exc


async def _read_pdf(path: Path, loop: asyncio.AbstractEventLoop) -> str:
    try:
        from pypdf import PdfReader  # type: ignore[import]
    except ImportError as exc:
        msg = "Reading .pdf files requires pypdf. Install with: pip install pypdf"
        raise PresentationError(msg) from exc

    def _extract() -> str:
        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                parts.append(text)
        return "\n\n".join(parts)

    try:
        return await loop.run_in_executor(None, _extract)
    except PresentationError:
        raise
    except Exception as exc:
        msg = f"Failed to parse {path.name} as a .pdf file: {exc}"
        raise PresentationError(msg) from exc


# ---------------------------------------------------------------------------
# 2a. Heuristic outline (no LLM required - always available)
# ---------------------------------------------------------------------------

_HEADING_RE = re.compile(r"^(#{1,3})\s+(.+)$", re.MULTILINE)


def _heuristic_outline(text: str, max_slides: int = _MAX_SLIDES_DEFAULT) -> list[SlidePlan]:
    """Split raw text into a slide outline without any LLM involved.

    Priority order:
    1. If the text has markdown headings (#, ##, ###), each heading
       starts a new slide and the text beneath it becomes bullets.
    2. Otherwise, split on blank-line-separated paragraphs (falling
       back to sentence splitting for a single dense block) and group
       them into up to `max_slides` roughly equal chunks, deriving a
       short title from each chunk's opening words.

    Always produces at least one slide for non-empty input and never
    raises - an effectively empty source simply yields no slides,
    which generate_presentation() reports as an error before any pptx
    machinery is touched.
    """
    text = text.strip()
    if not text:
        return []

    headings = list(_HEADING_RE.finditer(text))
    if headings:
        return _outline_from_headings(text, headings, max_slides)
    return _outline_from_paragraphs(text, max_slides)


def _outline_from_headings(
    text: str,
    headings: list[Match[str]],
    max_slides: int,
) -> list[SlidePlan]:
    slides: list[SlidePlan] = []
    for i, match in enumerate(headings):
        title = match.group(2).strip()
        body_start = match.end()
        body_end = headings[i + 1].start() if i + 1 < len(headings) else len(text)
        body = text[body_start:body_end].strip()
        bullets = _split_into_bullets(body)
        if title or bullets:
            slides.append(SlidePlan(title=title or f"Slide {i + 1}", bullets=bullets))
        if len(slides) >= max_slides:
            break
    return slides


def _outline_from_paragraphs(text: str, max_slides: int) -> list[SlidePlan]:
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    if not paragraphs:
        sentences = _split_sentences(text)
        paragraphs = sentences if sentences else [text]

    n_slides = min(max_slides, len(paragraphs)) or 1
    chunk_size = max(1, -(-len(paragraphs) // n_slides))  # ceil division
    slides: list[SlidePlan] = []
    for i in range(0, len(paragraphs), chunk_size):
        chunk = paragraphs[i : i + chunk_size]
        title = _derive_title(chunk[0])
        bullets: list[str] = []
        for para in chunk:
            bullets.extend(_split_into_bullets(para))
        slides.append(SlidePlan(title=title, bullets=bullets[:_MAX_BULLETS_PER_SLIDE]))
        if len(slides) >= max_slides:
            break
    return slides


def _split_into_bullets(body: str) -> list[str]:
    """Turn a block of text into a short list of bullet strings."""
    if not body:
        return []

    bullet_lines = re.findall(r"^\s*(?:[-*•]|\d+[.)])\s+(.+)$", body, re.MULTILINE)
    if bullet_lines:
        return [_truncate(b.strip()) for b in bullet_lines[:_MAX_BULLETS_PER_SLIDE]]

    paras = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
    if len(paras) > 1:
        return [_truncate(p) for p in paras[:_MAX_BULLETS_PER_SLIDE]]

    sentences = _split_sentences(body)
    return [_truncate(s) for s in sentences[:_MAX_BULLETS_PER_SLIDE]]


def _split_sentences(text: str) -> list[str]:
    text = " ".join(text.split())
    if not text:
        return []
    sentences = re.split(r"(?<=[.!?])\s+", text)
    return [s.strip() for s in sentences if s.strip()]


def _derive_title(paragraph: str) -> str:
    words = paragraph.split()
    title = " ".join(words[:8])
    return _truncate(title, max_chars=60) or "Untitled"


def _truncate(text: str, max_chars: int = _MAX_BULLET_CHARS) -> str:
    text = " ".join(text.split())
    if len(text) <= max_chars:
        return text
    return text[: max_chars - 1].rstrip() + "…"


# ---------------------------------------------------------------------------
# 2b. LLM-assisted outline (optional quality upgrade over the heuristic)
# ---------------------------------------------------------------------------

_LLM_OUTLINE_SYSTEM_PROMPT = (
    "You turn source documents into presentation outlines. Given the "
    "user's document, respond with ONLY a JSON array (no prose, no "
    "markdown fences) of objects shaped like "
    '{{"title": "...", "bullets": ["...", "..."]}}. '
    "Each slide needs a short title (under 8 words) and 2-5 concise "
    "bullets (each under 20 words). Produce at most {max_slides} "
    "slides. Cover the document's actual content - never invent facts "
    "that are not in the source."
)


async def _llm_outline(
    text: str,
    complete_fn: Any,
    max_slides: int = _MAX_SLIDES_DEFAULT,
    model: str | None = None,
) -> list[SlidePlan] | None:
    """Ask an LLM to build a richer outline than the heuristic splitter.

    `complete_fn` is expected to have the same shape as
    `LLMRuntime.complete()` used throughout
    integration/orchestrator.py's `_llm_complete()`: an async callable
    taking a single `CompletionRequest` and returning an object with a
    `.content` string. Returns None (never raises) on any failure -
    missing dependency, timeout, malformed JSON, empty result - so the
    caller can always fall back to `_heuristic_outline()`.
    """
    if complete_fn is None:
        return None

    try:
        from Agent_Guru.llm.models import CompletionRequest, Message
        from Agent_Guru.reasoning.llm.generation import ollama_model
    except ImportError:
        logger.debug("presentation: CompletionRequest/Message import failed")
        return None

    truncated_text = text[:_LLM_SOURCE_CHAR_LIMIT]
    system_prompt = _LLM_OUTLINE_SYSTEM_PROMPT.format(max_slides=max_slides)
    messages = [
        Message(role="system", content=system_prompt),
        Message(role="user", content=truncated_text),
    ]
    request = CompletionRequest(
        model=model or ollama_model(),
        messages=messages,
        max_tokens=1200,
    )

    try:
        response = await asyncio.wait_for(complete_fn(request), timeout=_LLM_OUTLINE_TIMEOUT_S)
    except TimeoutError:
        logger.debug("presentation: LLM outline timed out - using heuristic outline")
        return None
    except Exception as exc:
        logger.debug("presentation: LLM outline failed (%s) - using heuristic outline", exc)
        return None

    content = getattr(response, "content", None) or ""
    return _parse_llm_outline(content, max_slides)


def _parse_llm_outline(content: str, max_slides: int) -> list[SlidePlan] | None:
    content = content.strip()
    if content.startswith("```"):
        content = re.sub(r"^```(?:json)?\s*|\s*```$", "", content.strip(), flags=re.MULTILINE)

    try:
        data = json.loads(content)
    except (json.JSONDecodeError, ValueError):
        logger.debug("presentation: LLM outline was not valid JSON")
        return None

    if not isinstance(data, list) or not data:
        return None

    slides: list[SlidePlan] = []
    for item in data[:max_slides]:
        if not isinstance(item, dict):
            continue
        title = str(item.get("title", "")).strip()
        raw_bullets = item.get("bullets", [])
        if not isinstance(raw_bullets, list):
            continue
        bullets = [
            _truncate(str(b).strip())
            for b in raw_bullets[:_MAX_BULLETS_PER_SLIDE]
            if str(b).strip()
        ]
        if title and bullets:
            slides.append(SlidePlan(title=_truncate(title, max_chars=60), bullets=bullets))

    return slides or None


# ---------------------------------------------------------------------------
# 3. Rendering to a real .pptx
# ---------------------------------------------------------------------------


def render_pptx(
    outline: list[SlidePlan],
    output_path: Path,
    deck_title: str | None = None,
) -> Path:
    """Render a slide outline into a real .pptx file at output_path.

    Synchronous - python-pptx has no async API. generate_presentation()
    below runs this in an executor so it never blocks the event loop.
    """
    try:
        from pptx import Presentation  # type: ignore[import]
    except ImportError as exc:
        msg = "Rendering .pptx files requires python-pptx. Install with: pip install python-pptx"
        raise PresentationError(msg) from exc

    if not outline:
        msg = "Cannot render a presentation with an empty outline"
        raise PresentationError(msg)

    prs = Presentation()

    if deck_title:
        title_layout = prs.slide_layouts[0]  # "Title Slide"
        title_slide = prs.slides.add_slide(title_layout)
        title_slide.shapes.title.text = deck_title
        if len(title_slide.placeholders) > 1:
            title_slide.placeholders[1].text = "Generated by Agent Guru"

    content_layout = prs.slide_layouts[1]  # "Title and Content"
    for slide_plan in outline:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = slide_plan.title

        body_placeholder = slide.placeholders[1]
        text_frame = body_placeholder.text_frame
        text_frame.clear()
        for i, bullet in enumerate(slide_plan.bullets):
            paragraph = text_frame.paragraphs[0] if i == 0 else text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    return output_path


# ---------------------------------------------------------------------------
# Public entry point (local, offline-first path)
# ---------------------------------------------------------------------------


async def generate_presentation(
    source_path: str | Path,
    output_path: str | Path | None = None,
    *,
    complete_fn: Any = None,
    llm_model: str | None = None,
    deck_title: str | None = None,
    max_slides: int = _MAX_SLIDES_DEFAULT,
) -> Path:
    """Generate a .pptx from a source document - this module's main entry point.

    Args:
        source_path: Path to the source file (.txt, .md, .docx, .pdf,
            or any other text-readable format).
        output_path: Where to save the .pptx. Defaults to the source
            file's own directory, same stem, ".pptx" extension.
        complete_fn: Optional async LLM completion callable (same
            shape as `LLMRuntime.complete()`) for a richer outline.
            When None (the default) or when it fails for any reason,
            the dependency-free heuristic outline is used instead -
            this function ALWAYS succeeds at producing *a* deck for
            any non-empty source, LLM or not.
        llm_model: Model name to request if complete_fn is given.
        deck_title: Title slide text. Defaults to the source
            filename (underscores/dashes replaced, title-cased).
        max_slides: Upper bound on slide count.

    Returns:
        Path to the generated .pptx file.

    Raises:
        PresentationError: source file missing/unreadable/empty, or
            python-pptx (or, for .docx/.pdf sources, python-docx/pypdf)
            is not installed.

    """
    source_path = Path(source_path)
    if not source_path.is_file():
        msg = f"Source file not found: {source_path}"
        raise PresentationError(msg)

    resolved_output = Path(output_path) if output_path is not None else source_path.with_suffix(".pptx")

    text = await _read_source_text(source_path)
    if not text or not text.strip():
        msg = f"Source file has no extractable text: {source_path}"
        raise PresentationError(msg)

    outline: list[SlidePlan] | None = None
    if complete_fn is not None:
        outline = await _llm_outline(text, complete_fn, max_slides, model=llm_model)
        if outline:
            logger.info("presentation: using LLM-assisted outline (%d slides)", len(outline))

    if not outline:
        outline = _heuristic_outline(text, max_slides)
        logger.info("presentation: using heuristic outline (%d slides)", len(outline))

    if not outline:
        msg = f"Could not build any slides from {source_path.name} - the file appears to have no usable text content"
        raise PresentationError(msg)

    title = deck_title or source_path.stem.replace("_", " ").replace("-", " ").strip().title()

    loop = asyncio.get_running_loop()
    return await loop.run_in_executor(None, render_pptx, outline, resolved_output, title)


# ---------------------------------------------------------------------------
# Optional cloud path: Gamma API (opt-in only, never the default)
# ---------------------------------------------------------------------------

_GAMMA_ENDPOINT = "https://public-api.gamma.app/v1.0/generations"
_GAMMA_MAX_CHARS = 400_000


async def generate_presentation_gamma(
    source_path: str | Path,
    *,
    api_key: str | None = None,
    title: str | None = None,
    num_cards: int | None = None,
    additional_instructions: str | None = None,
) -> dict[str, Any]:
    """Generate a presentation via Gamma's cloud API instead of locally.

    Entirely optional and OFF by default - Agent Guru is offline-first
    and `generate_presentation()` above is the local path that always
    works with no network and no account. Use this only when the user
    has explicitly asked for a Gamma-generated deck and has their own
    Gamma API key.

    api_key is never read implicitly from anywhere the user did not
    put it themselves: pass it explicitly, or set the GAMMA_API_KEY
    environment variable yourself before calling this - there is no
    silent fallback between the two, and if neither is present this
    raises rather than guessing or sending data without a key.

    Gamma's generation endpoint is asynchronous: this call starts a
    generation job and returns Gamma's response payload (containing a
    generationId to poll for the finished file), not a finished file
    itself - see
    https://developers.gamma.app/generations/create-generation for the
    full, evolving contract.

    Raises:
        PresentationError: no API key available, source unreadable,
            aiohttp missing, or the HTTP request fails.

    """
    resolved_key = api_key or os.environ.get("GAMMA_API_KEY")
    if not resolved_key:
        msg = (
            "Gamma generation requires an API key. Pass api_key= "
            "explicitly or set the GAMMA_API_KEY environment variable "
            "- Agent Guru never sends your document to Gamma without "
            "an explicit key you provided."
        )
        raise PresentationError(msg)

    source_path = Path(source_path)
    if not source_path.is_file():
        msg = f"Source file not found: {source_path}"
        raise PresentationError(msg)

    text = await _read_source_text(source_path)
    if not text or not text.strip():
        msg = f"Source file has no extractable text: {source_path}"
        raise PresentationError(msg)
    text = text[:_GAMMA_MAX_CHARS]

    try:
        import aiohttp
    except ImportError as exc:
        msg = "Gamma generation requires aiohttp. Install with: pip install aiohttp"
        raise PresentationError(msg) from exc

    payload: dict[str, Any] = {
        "inputText": text,
        "format": "presentation",
        "textMode": "condense",
    }
    if title:
        payload["title"] = title
    if num_cards:
        payload["numCards"] = num_cards
    if additional_instructions:
        payload["additionalInstructions"] = additional_instructions

    headers = {"X-API-KEY": resolved_key, "Content-Type": "application/json"}

    try:
        async with (
            aiohttp.ClientSession() as session,
            session.post(
                _GAMMA_ENDPOINT,
                json=payload,
                headers=headers,
                timeout=aiohttp.ClientTimeout(total=30),
            ) as resp,
        ):
            body = await resp.json(content_type=None)
            if resp.status >= 400:
                msg = f"Gamma API returned {resp.status}: {body}"
                raise PresentationError(msg)
            return body
    except PresentationError:
        raise
    except TimeoutError as exc:
        msg = "Gamma API request timed out"
        raise PresentationError(msg) from exc
    except Exception as exc:
        msg = f"Gamma API request failed: {exc}"
        raise PresentationError(msg) from exc
